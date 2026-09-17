#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
用途：
1. 读取 Excel 中的“停车系统订单ID”和“订单金额”
2. 查询 MySQL 表 road_bill（type=1）
3. 逐笔比对 Excel 金额 和 road_bill.amount
4. 输出带比对结果的新 Excel，并生成汇总 sheet

依赖：
    pip install pandas openpyxl pymysql

运行：
    python check_road_bill_excel.py \
      --host 127.0.0.1 --port 3306 --user root --password 123456 --database your_db \
      --input "/path/to/副本智城1月份助缴订单明细.xlsx" \
      --output "/path/to/副本智城1月份助缴订单明细_比对结果.xlsx"

说明：
- 默认只查 road_bill.type = 1
- 默认过滤 is_deleted = 0（如果你的表确实有这个字段，且逻辑删除不需要参与比对）
- 若同一个 road_order_id 在 DB 中查到多笔 type=1 记录，会标注为“DB存在多笔支付账单”
- 若 Excel 中订单号为空/非法，也会标注
"""

from __future__ import annotations

import argparse
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Dict, List, Any

import pandas as pd
import pymysql
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

EXCEL_ORDER_ID_COL = "停车系统订单ID"
EXCEL_AMOUNT_COL = "订单金额"
EXCEL_PAY_TIME_COL = "交易成功时间"

RESULT_COLS = [
    "DB账单数量",
    "DB账单ID列表",
    "DB金额列表",
    "DB金额合计",
    "DB支付方式",
    "比对结果",
    "异常说明",
]

OK_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")  # 浅绿
WARN_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")  # 浅黄
ERR_FILL = PatternFill(fill_type="solid", fgColor="F4CCCC")  # 浅红


def to_decimal(val: Any) -> Decimal | None:
    if val is None:
        return None
    s = str(val).strip()
    if s == "" or s.lower() == "nan":
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def normalize_order_id(val: Any) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s


def chunked(seq: List[str], size: int):
    for i in range(0, len(seq), size):
        yield seq[i:i + size]


def fetch_bill_map(conn, order_ids: List[str], table_name: str, use_is_deleted: bool = True) -> Dict[str, List[dict]]:
    """返回：{road_order_id: [row, row, ...]}"""
    result: Dict[str, List[dict]] = {}
    if not order_ids:
        return result

    base_sql = f"""
        SELECT
            id,
            road_order_id,
            amount,
            type,
            status,
            create_time,
            success_time,
            trade_no,
            is_deleted,
            pay_mode,
            property_attribution_id
        FROM {table_name}
        WHERE type = 1
          AND road_order_id IN (%s)
    """
    if use_is_deleted:
        base_sql += " AND is_deleted = 0"
    base_sql += " ORDER BY road_order_id, id"

    with conn.cursor(pymysql.cursors.DictCursor) as cursor:
        for batch in chunked(order_ids, 1000):
            placeholders = ",".join(["%s"] * len(batch))
            sql = base_sql.replace("%s", placeholders, 1)
            cursor.execute(sql, batch)
            rows = cursor.fetchall()
            for row in rows:
                key = str(row["road_order_id"])
                result.setdefault(key, []).append(row)
    return result


def build_compare_result(df: pd.DataFrame, bill_map: Dict[str, List[dict]]) -> pd.DataFrame:
    db_count_list = []
    db_id_list = []
    db_amount_list = []
    db_amount_sum_list = []
    db_pay_mode_list = []
    compare_result_list = []
    remark_list = []

    for _, row in df.iterrows():
        order_id = normalize_order_id(row.get(EXCEL_ORDER_ID_COL))
        excel_amount = to_decimal(row.get(EXCEL_AMOUNT_COL))
        excel_pay_time = row.get(EXCEL_PAY_TIME_COL)

        matched_rows = bill_map.get(order_id, []) if order_id else []
        count = len(matched_rows)
        amounts = [to_decimal(x.get("amount")) for x in matched_rows]
        amount_sum = sum((a for a in amounts if a is not None), Decimal("0.00")) if amounts else None

        db_count_list.append(count)
        db_id_list.append(",".join(str(x["id"]) for x in matched_rows) if matched_rows else "")
        db_amount_list.append(",".join(str(x["amount"]) for x in matched_rows) if matched_rows else "")
        db_amount_sum_list.append(str(amount_sum) if amount_sum is not None else "")
        db_pay_mode_list.append(",".join(str(x.get("pay_mode", "")) for x in matched_rows) if matched_rows else "")

        if not order_id:
            compare_result_list.append("异常")
            remark_list.append("Excel订单ID为空")
            continue

        if excel_amount is None:
            compare_result_list.append("异常")
            remark_list.append("Excel订单金额为空或格式非法")
            continue

        if count == 0:
            compare_result_list.append("异常")
            remark_list.append("数据库查不到该订单(type=1)")
            continue

        if count > 1:
            # 多笔支付账单先直接标异常，同时保留金额合计供人工判断
            compare_result_list.append("异常")
            if amount_sum == excel_amount:
                remark_list.append("DB存在多笔支付账单，但金额合计与Excel一致")
            else:
                remark_list.append("DB存在多笔支付账单，且金额合计与Excel不一致")
            continue

        db_amount = amounts[0]
        db_pay_mode = matched_rows[0].get("pay_mode")
        db_create_time = matched_rows[0].get("create_time")
        db_property_attribution_id = matched_rows[0].get("property_attribution_id")

        if db_amount is None:
            compare_result_list.append("异常")
            remark_list.append("数据库amount为空或格式非法")
            continue

        if db_amount != excel_amount:
            compare_result_list.append("异常")
            remark_list.append(f"金额不一致：Excel={excel_amount}，DB={db_amount}")
            continue

        if str(db_pay_mode) != "23":
            compare_result_list.append("异常")
            remark_list.append(f"支付方式异常：预期23，实际为{db_pay_mode}")
            continue

        if db_create_time and excel_pay_time and pd.notna(excel_pay_time):
            try:
                # 尝试解析 excel_pay_time
                excel_date = pd.to_datetime(excel_pay_time)
                # create_time 通常是 datetime 对象或字符串
                db_date = pd.to_datetime(db_create_time)
                if excel_date.month != db_date.month or excel_date.year != db_date.year:
                    compare_result_list.append("异常")
                    remark_list.append(
                        f"月份不一致：Excel交易时间={excel_date.strftime('%Y-%m')}，DB创建时间={db_date.strftime('%Y-%m')}")
                    continue
            except Exception:
                # 日期解析失败的情况跳过判断或做特殊处理
                pass

        valid_property_ids = {"1920040795040256000", "1948648310693986304"}
        if str(db_property_attribution_id) not in valid_property_ids:
            compare_result_list.append("异常")
            remark_list.append(f"归属异常：实际为{db_property_attribution_id}")
            continue

        compare_result_list.append("匹配")
        remark_list.append("")

    df["DB账单数量"] = db_count_list
    df["DB账单ID列表"] = db_id_list
    df["DB金额列表"] = db_amount_list
    df["DB金额合计"] = db_amount_sum_list
    df["DB支付方式"] = db_pay_mode_list
    df["比对结果"] = compare_result_list
    df["异常说明"] = remark_list
    return df


def auto_adjust_excel(output_file: Path):
    wb = load_workbook(output_file)
    ws = wb["比对结果"]

    header_map = {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}
    result_col = header_map.get("比对结果")

    # 自动列宽
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 40)

    # 结果着色
    if result_col:
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=result_col)
            if cell.value == "匹配":
                for c in ws[row]:
                    c.fill = OK_FILL
            elif cell.value == "异常":
                for c in ws[row]:
                    c.fill = ERR_FILL
            else:
                for c in ws[row]:
                    c.fill = WARN_FILL

    wb.save(output_file)


def check_table_has_is_deleted(conn, table_name: str) -> bool:
    sql = """
        SELECT COUNT(*) AS cnt
        FROM information_schema.COLUMNS
        WHERE TABLE_SCHEMA = DATABASE()
          AND TABLE_NAME = %s
          AND COLUMN_NAME = 'is_deleted'
    """
    with conn.cursor(pymysql.cursors.DictCursor) as cursor:
        cursor.execute(sql, (table_name,))
        row = cursor.fetchone()
        return bool(row and row["cnt"] > 0)


def main():
    parser = argparse.ArgumentParser(description="Excel订单金额 vs road_bill.amount 对账脚本")
    parser.add_argument("--host", required=True, help="MySQL host")
    parser.add_argument("--port", type=int, default=3306, help="MySQL port")
    parser.add_argument("--user", required=True, help="MySQL user")
    parser.add_argument("--password", required=True, help="MySQL password")
    parser.add_argument("--database", required=True, help="MySQL database")
    parser.add_argument("--input", required=True, help="输入Excel路径")
    parser.add_argument("--output", required=True, help="输出Excel路径")
    parser.add_argument("--table", default="road_bill", help="表名，默认 road_bill")
    parser.add_argument("--sheet", default=0, help="sheet名或索引，默认第一个sheet")
    args = parser.parse_args()

    print("3")
    input_file = Path(args.input)
    output_file = Path(args.output)
    print("4")
    if not input_file.exists():
        raise FileNotFoundError(f"输入文件不存在: {input_file}")

    print("1")
    df = pd.read_excel(input_file, sheet_name=args.sheet, dtype=str)
    print("2")
    if EXCEL_ORDER_ID_COL not in df.columns:
        raise ValueError(f"Excel缺少列：{EXCEL_ORDER_ID_COL}")
    if EXCEL_AMOUNT_COL not in df.columns:
        raise ValueError(f"Excel缺少列：{EXCEL_AMOUNT_COL}")
    if EXCEL_PAY_TIME_COL not in df.columns:
        raise ValueError(f"Excel缺少列：{EXCEL_PAY_TIME_COL}")

    df[EXCEL_ORDER_ID_COL] = df[EXCEL_ORDER_ID_COL].map(normalize_order_id)

    unique_order_ids = sorted({x for x in df[EXCEL_ORDER_ID_COL].tolist() if x})
    print(f"Excel总行数: {len(df)}")
    print(f"待查询唯一订单数: {len(unique_order_ids)}")

    conn = pymysql.connect(
        host=args.host,
        port=args.port,
        user=args.user,
        password=args.password,
        database=args.database,
        charset="utf8mb4",
        autocommit=True,
    )
    try:
        use_is_deleted = check_table_has_is_deleted(conn, args.table)
        bill_map = fetch_bill_map(conn, unique_order_ids, args.table, use_is_deleted=use_is_deleted)
        result_df = build_compare_result(df, bill_map)

        summary = {
            "Excel总行数": len(result_df),
            "匹配笔数": int((result_df["比对结果"] == "匹配").sum()),
            "异常笔数": int((result_df["比对结果"] == "异常").sum()),
            "Excel金额总计": str(
                sum((to_decimal(x) or Decimal("0.00") for x in result_df[EXCEL_AMOUNT_COL]), Decimal("0.00"))),
            "异常金额总计": str(sum((to_decimal(row[EXCEL_AMOUNT_COL]) or Decimal("0.00")
                                     for _, row in result_df.iterrows()
                                     if row["比对结果"] == "异常"), Decimal("0.00"))),
            "DB金额合计(仅按结果表汇总列求和)": str(
                sum((to_decimal(x) or Decimal("0.00") for x in result_df["DB金额合计"]), Decimal("0.00"))),
            "金额差额(Excel-DB)": str(
                sum((to_decimal(x) or Decimal("0.00") for x in result_df[EXCEL_AMOUNT_COL]), Decimal("0.00"))
                - sum((to_decimal(x) or Decimal("0.00") for x in result_df["DB金额合计"]), Decimal("0.00"))
            ),
        }

        exception_detail = (
            result_df[result_df["比对结果"] == "异常"]
            .groupby("异常说明", dropna=False)
            .size()
            .reset_index(name="数量")
            .sort_values("数量", ascending=False)
        )

        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            result_df.to_excel(writer, sheet_name="比对结果", index=False)
            pd.DataFrame([summary]).to_excel(writer, sheet_name="汇总", index=False)
            exception_detail.to_excel(writer, sheet_name="异常分类统计", index=False)

        auto_adjust_excel(output_file)

        print("\n比对完成")
        print(f"输出文件: {output_file}")
        for k, v in summary.items():
            print(f"{k}: {v}")

    finally:
        conn.close()


if __name__ == "__main__":
    main()
